from docx import Document

d = Document(r'C:\Users\Kishore GK\Desktop\VIDS\Report\Voice_Impersonation_Detection_System_Report.docx')
paras = [p.text for p in d.paragraphs]
print('paragraphs:', len(paras), '| tables:', len(d.tables))
print('--- placeholder hunt ---')
found = 0
for t in paras:
    if 'To be updated' in t or 'TODO' in t or 'next phase' in t.lower():
        found += 1
        print('FOUND:', t[:110])
print('placeholders:', found)
print('--- chapter 7 extract ---')
in7 = False
for t in paras:
    if 'CHAPTER 7' in t:
        in7 = True
    if in7:
        print(repr(t[:130]))
        if 'CHAPTER 8' in t:
            break
print('--- inline images ---')
print('images:', len(d.inline_shapes))
