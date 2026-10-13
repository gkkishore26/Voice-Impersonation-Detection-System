import re
from docx import Document

d = Document(r'C:\Users\Kishore GK\Desktop\VIDS\Report\Voice_Impersonation_Detection_System_Report.docx')
full = '\n'.join(p.text for p in d.paragraphs)
print('tamil script chars:', len(re.findall(r'[\u0B80-\u0BFF]', full)))
print('has KISHORE:', 'KISHORE' in full)
print('has Karpagam:', 'Karpagam' in full)
print('has 24ADU297:', '24ADU297' in full)
print('has Professor K:', 'Professor K' in full)
for t in d.paragraphs:
    if 'KISHORE' in t.text or 'Karpagam Academy' in t.text:
        print('TITLE:', t.text[:120])
        break
tables_text = ' | '.join(c.text for t in d.tables for r in t.rows for c in r.cells)
import re as _re
m = _re.search(r'Accuracy.{0,40}9\d\.\d%', tables_text)
print('table accuracy cell:', m.group(0) if m else 'check manually')
